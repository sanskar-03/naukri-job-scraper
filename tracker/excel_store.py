from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import re

HEADERS = [
    "job_id",
    "title",
    "company",
    "location",
    "experience",
    "skills",
    "posted_date",
    "job_url",
    "scraped_at",
]

MASTER_SHEET = "Master History"


def clean_sheet_name(keyword, location):
    name = f"{keyword} - {location}".strip()

    # Excel-invalid characters
    name = re.sub(r'[\\/*?:\[\]]', "-", name)

    # Collapse spaces
    name = re.sub(r"\s+", " ", name).strip()

    if not name:
        name = "Search"

    # Excel maximum worksheet length = 31
    return name[:31]


def ensure_workbook(path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if not path.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = MASTER_SHEET

        ws.append(HEADERS)

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid",
                fgColor="1f2937"
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        wb.save(path)
        wb.close()


def prepare_sheet(ws):
    ws.freeze_panes = "A2"

    if ws.max_row >= 1:
        ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 22,
        "B": 42,
        "C": 32,
        "D": 32,
        "E": 20,
        "F": 65,
        "G": 20,
        "H": 85,
        "I": 24,
    }

    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def existing_keys(ws):
    ids = set()
    urls = set()

    header_map = {}

    for index, cell in enumerate(ws[1], start=1):
        if cell.value:
            header_map[str(cell.value).strip()] = index

    id_col = header_map.get("job_id")
    url_col = header_map.get("job_url")

    if not id_col:
        return ids, urls

    for row in ws.iter_rows(min_row=2, values_only=True):

        if id_col <= len(row):
            value = row[id_col - 1]

            if value:
                ids.add(str(value).strip())

        if url_col and url_col <= len(row):
            value = row[url_col - 1]

            if value:
                urls.add(str(value).strip())

    return ids, urls


def get_or_create_search_sheet(wb, keyword, location):
    sheet_name = clean_sheet_name(keyword, location)

    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
        ws.append(HEADERS)

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid",
                fgColor="1f2937"
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

    prepare_sheet(ws)

    return ws


def read_sheet_jobs(ws):
    jobs = []

    headers = [
        str(cell.value).strip()
        if cell.value is not None
        else ""
        for cell in ws[1]
    ]

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):
        if not any(
            value not in (None, "")
            for value in row
        ):
            continue

        job = {}

        for index, header in enumerate(headers):
            if header:
                job[header] = (
                    row[index]
                    if index < len(row)
                    else ""
                )

        # Filter out jobs that are 30+ days old (expired)
        posted = str(job.get("posted_date", "")).lower()
        if "30+" in posted or "30 days" in posted or "expired" in posted:
            continue

        jobs.append(job)

    return jobs




def read_jobs(path):
    ensure_workbook(path)

    wb = load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    jobs = []

    for sheet_name in wb.sheetnames:

        if sheet_name == MASTER_SHEET:
            continue

        ws = wb[sheet_name]

        jobs.extend(
            read_sheet_jobs(ws)
        )

    wb.close()

    return jobs


def read_search_jobs(path, keyword, location):
    ensure_workbook(path)

    wb = load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    sheet_name = clean_sheet_name(
        keyword,
        location
    )

    if sheet_name not in wb.sheetnames:
        wb.close()
        return []

    ws = wb[sheet_name]

    jobs = read_sheet_jobs(ws)

    wb.close()

    return jobs


def append_new_jobs(
    path,
    jobs,
    keyword,
    location
):
    ensure_workbook(path)

    wb = load_workbook(path)

    search_ws = get_or_create_search_sheet(
        wb,
        keyword,
        location
    )

    master_ws = wb[MASTER_SHEET]

    search_ids, search_urls = existing_keys(
        search_ws
    )

    master_ids, master_urls = existing_keys(
        master_ws
    )

    new_jobs = []

    for job in jobs:

        job_id = str(
            job.get("job_id") or ""
        ).strip()

        job_url = str(
            job.get("job_url") or ""
        ).strip()

        if not job_id and not job_url:
            continue

        # ----------------------------------------------------
        # DUPLICATE INSIDE THIS SEARCH
        # ----------------------------------------------------

        if (
            (job_id and job_id in search_ids)
            or
            (job_url and job_url in search_urls)
        ):
            continue

        row = [
            job.get(header, "")
            for header in HEADERS
        ]

        # ----------------------------------------------------
        # SEARCH SHEET
        # ----------------------------------------------------

        search_ws.append(row)

        if job_id:
            search_ids.add(job_id)

        if job_url:
            search_urls.add(job_url)

        new_jobs.append(job)

        # ----------------------------------------------------
        # MASTER HISTORY
        # ----------------------------------------------------

        if (
            (job_id and job_id in master_ids)
            or
            (job_url and job_url in master_urls)
        ):
            continue

        master_ws.append(row)

        if job_id:
            master_ids.add(job_id)

        if job_url:
            master_urls.add(job_url)

    prepare_sheet(search_ws)
    prepare_sheet(master_ws)

    wb.save(path)
    wb.close()

    return new_jobs


def export_search_excel(
    path,
    keyword,
    location,
    export_dir
):
    export_dir = Path(export_dir)
    export_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    jobs = read_search_jobs(
        path,
        keyword,
        location
    )

    safe_keyword = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        str(keyword).strip()
    ).strip("_")

    safe_location = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        str(location).strip()
    ).strip("_")

    filename = (
        f"{safe_keyword}_{safe_location}.xlsx"
    )

    output = export_dir / filename

    wb = Workbook()
    ws = wb.active
    ws.title = "Job Results"

    ws.append(HEADERS)

    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            "solid",
            fgColor="1f2937"
        )
        cell.alignment = Alignment(
            horizontal="center"
        )

    for job in jobs:
        ws.append([
            job.get(header, "")
            for header in HEADERS
        ])

    prepare_sheet(ws)

    wb.save(output)
    wb.close()

    return output


def get_search_history(path):
    ensure_workbook(path)

    wb = load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    history = []

    for sheet_name in wb.sheetnames:

        if sheet_name == MASTER_SHEET:
            continue

        ws = wb[sheet_name]

        jobs = read_sheet_jobs(ws)

        history.append({
            "search": sheet_name,
            "count": len(jobs),
        })

    wb.close()

    return history


def delete_search(path, keyword="", location="", search_name=None):
    ensure_workbook(path)
    wb = load_workbook(path)

    target_sheet = search_name
    if not target_sheet and (keyword or location):
        target_sheet = clean_sheet_name(keyword, location)

    if not target_sheet:
        wb.close()
        return False

    found_sheet = None
    for name in wb.sheetnames:
        if name.lower() == str(target_sheet).lower():
            found_sheet = name
            break

    if not found_sheet or found_sheet == MASTER_SHEET:
        wb.close()
        return False

    del wb[found_sheet]

    # Rebuild Master History from remaining sheets so it remains consistent
    if MASTER_SHEET in wb.sheetnames:
        del wb[MASTER_SHEET]

    master_ws = wb.create_sheet(MASTER_SHEET, 0)
    master_ws.append(HEADERS)
    for cell in master_ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1f2937")
        cell.alignment = Alignment(horizontal="center")
    prepare_sheet(master_ws)

    seen_keys = set()
    for name in wb.sheetnames:
        if name == MASTER_SHEET:
            continue
        ws = wb[name]
        for job in read_sheet_jobs(ws):
            key = str(job.get("job_id") or "") + "|" + str(job.get("job_url") or "")
            if key not in seen_keys:
                seen_keys.add(key)
                master_ws.append([job.get(h, "") for h in HEADERS])

    wb.save(path)
    wb.close()
    return True


def delete_all_searches(path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = MASTER_SHEET
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1f2937")
        cell.alignment = Alignment(horizontal="center")
    prepare_sheet(ws)

    wb.save(path)
    wb.close()
    return True

